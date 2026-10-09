"""تصدير التقارير إلى Excel و CSV."""

import csv
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

MONEY = '#,##0.00'
INT = '#,##0'
PCT = '0.0"%"'

ORDER_COLUMNS = [
    ("رقم الطلب", "id", None),
    ("التاريخ", "date", None),
    ("المتجر", "store_name", None),
    ("العميل", "customer", None),
    ("المدينة", "city", None),
    ("طريقة الدفع", "payment", None),
    ("الحالة", "status", None),
    ("عدد القطع", "qty", INT),
    ("المنتجات قبل الخصم", "subtotal", MONEY),
    ("الخصم", "discount", MONEY),
    ("الشحن المحصّل", "shipping", MONEY),
    ("الإجمالي شامل الضريبة", "total", MONEY),
    ("ضريبة القيمة المضافة", "vat", MONEY),
    ("الصافي بدون ضريبة", "net", MONEY),
    ("تكلفة البضاعة", "cogs", MONEY),
    ("تكلفة الشحن", "ship_cost", MONEY),
    ("رسوم الدفع", "pay_fee", MONEY),
    ("الربح التقديري", "profit", MONEY),
]


def _order_row(o, store_names):
    row = dict(o, store_name=store_names.get(o["store"], o["store"]))
    if o["status"] == "ملغي":
        # الطلب الملغي ما بيدخلش في الإيراد
        for k in ("vat", "net", "cogs", "ship_cost", "pay_fee", "profit"):
            row[k] = 0
    return row


def _style_header(ws, row, fill):
    for cell in ws[row]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _autosize(ws, min_w=10, max_w=42):
    for col in ws.columns:
        width = max((len(str(c.value)) for c in col if c.value is not None), default=0)
        ws.column_dimensions[get_column_letter(col[0].column)].width = max(min_w, min(max_w, width + 4))


def _table_sheet(wb, title, headers, rows, formats, fill):
    ws = wb.create_sheet(title)
    ws.sheet_view.rightToLeft = True
    ws.append(headers)
    _style_header(ws, 1, fill)
    for r in rows:
        ws.append(r)
    for idx, fmt in enumerate(formats, start=1):
        if fmt:
            for cell in ws.iter_cols(min_col=idx, max_col=idx, min_row=2):
                for c in cell:
                    c.number_format = fmt
    ws.freeze_panes = "A2"
    _autosize(ws)
    return ws


def build_xlsx(report, store_label, store_names, brand_hex):
    fill = PatternFill("solid", fgColor=brand_hex.lstrip("#").upper())
    soft = PatternFill("solid", fgColor="F4F4F8")
    thin = Side(style="thin", color="DDDDE6")
    m, period = report["m"], report["period"]

    wb = Workbook()
    ws = wb.active
    ws.title = "الملخص"
    ws.sheet_view.rightToLeft = True
    ws.append([f"تقرير {report['kinds'][report['kind']]} — {store_label}"])
    ws["A1"].font = Font(bold=True, size=15)
    ws.append([f"الفترة: {period['label']}" + (" (فترة غير مكتملة)" if period["partial"] else "")])
    ws.append(["بيانات تجريبية للعرض"])
    ws["A3"].font = Font(italic=True, color="A75A06")
    ws.append([])

    ws.append(["قائمة الدخل", "المبلغ (ر.س)"])
    _style_header(ws, ws.max_row, fill)
    for line in report["statement"]:
        ws.append([line["label"], round(line["value"], 2)])
        c = ws.cell(row=ws.max_row, column=2)
        c.number_format = MONEY
        if line["kind"] in ("subtotal", "result", "total"):
            for cell in ws[ws.max_row]:
                cell.font = Font(bold=True)
                cell.fill = soft
        for cell in ws[ws.max_row]:
            cell.border = Border(bottom=thin)
    ws.append([])

    ws.append(["المؤشر", "القيمة"])
    _style_header(ws, ws.max_row, fill)
    kpis = [
        ("عدد الطلبات المكتملة", m["orders"], INT),
        ("الطلبات الملغية", m["cancelled"], INT),
        ("نسبة الإلغاء", m["cancel_rate"], PCT),
        ("متوسط قيمة الطلب", m["aov"], MONEY),
        ("متوسط القطع في الطلب", m["items_per_order"], '0.00'),
        ("إجمالي الخصومات", m["discount"], MONEY),
        ("نسبة الخصم من المبيعات", m["discount_rate"], PCT),
        ("الشحن المحصّل من العملاء (شامل الضريبة)", m["shipping"], MONEY),
        ("صافي الشحن (الإيراد - التكلفة)", m["ship_net"], MONEY),
        ("طلبات بشحن مجاني", m["free_ship"], INT),
        ("مجمل الربح من المنتجات", m["gross_profit"], MONEY),
        ("هامش صافي الربح", m["margin"], PCT),
        ("عدد العملاء", m["customers"], INT),
        ("عملاء جدد", m["new_customers"], INT),
        ("عملاء عائدون", m["returning_customers"], INT),
    ]
    for label, value, fmt in kpis:
        ws.append([label, round(value, 2)])
        ws.cell(row=ws.max_row, column=2).number_format = fmt
    ws.column_dimensions["A"].width = 46
    ws.column_dimensions["B"].width = 20

    orders = [_order_row(o, store_names) for o in report["orders"]]
    _table_sheet(wb, "الطلبات", [c[0] for c in ORDER_COLUMNS],
                 [[round(o[c[1]], 2) if isinstance(o[c[1]], float) else o[c[1]] for c in ORDER_COLUMNS] for o in orders],
                 [c[2] for c in ORDER_COLUMNS], fill)

    _table_sheet(wb, "المنتجات", ["المنتج", "المتجر", "القطع المباعة", "عدد الطلبات", "المبيعات شامل الضريبة", "مجمل الربح", "الهامش"],
                 [[p["name"], store_names[p["store"]], p["qty"], p["orders"], round(p["sales"], 2), round(p["profit"], 2), round(p["margin"], 1)]
                  for p in report["products"]],
                 [None, None, INT, INT, MONEY, MONEY, PCT], fill)

    for title, rows in (("المدن", report["cities"]), ("طرق الدفع", report["payments"])):
        _table_sheet(wb, title, ["الاسم", "عدد الطلبات", "المبيعات", "النسبة"],
                     [[r["name"], r["orders"], round(r["sales"], 2), round(r["share"], 1)] for r in rows],
                     [None, INT, MONEY, PCT], fill)

    _table_sheet(wb, "الاتجاه", ["الفترة"] + [store_names[s] for s in ("smooth", "glorias") if s in report["trend"][0]],
                 [[t["full"]] + [t[s] for s in ("smooth", "glorias") if s in t] for t in report["trend"]],
                 [None, MONEY, MONEY], fill)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_csv(report, store_names):
    buf = io.StringIO()
    buf.write("﻿")  # عشان Excel يقرا العربي صح
    w = csv.writer(buf)
    w.writerow([c[0] for c in ORDER_COLUMNS])
    for o in report["orders"]:
        row = _order_row(o, store_names)
        w.writerow([round(row[c[1]], 2) if isinstance(row[c[1]], float) else row[c[1]] for c in ORDER_COLUMNS])
    return buf.getvalue().encode("utf-8")
